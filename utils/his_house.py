from functools import reduce
from json import dump
from pprint import pprint
from docx import Document, document
from re import MULTILINE, sub, findall

from docx.text.parfmt import ParagraphFormat
from docx.shared import Pt
doc = Document('utils/His House Worship Songs 1.3.docx')

def get_title(count, content) -> bool:
    count = int(count)
    if content == '':
        count += 1
    else:
        # Rest if not continous
        count = 0
    if count > 4:
        count = str(count)
        return True
    else:
        count = str(count)
        return False
# titles = list(
#     reduce(
#     lambda count, content: content if get_title(count,content) else count,
#     all_text,
#     [0]) # Is the first var
#     )

# print(titles)

def find_songs():
    all_text: list[str] = list(map(lambda x: x.text, doc.paragraphs))
    count = 0
    title = False
    titles = []
    for line in all_text:
        if '161' in line:
            ... # Line to catch the debugger
        if line == '':
            count += 1
        elif not title:
            count = 0
        if count > 4:
            title = True
        if count > 4 and title and line != '':
            print(line)
            title = False
            titles.append(line)
    return titles

# titles = find_songs(all_text) # So far missing a few

# print(f"Len of titles is: {len(titles)}")

def bold_titles(doc):
    for paragrapgh in doc.paragraphs:
        curr_run = ""
        for run in paragrapgh.runs:
            if run.bold:
                curr_run += run.text
    # print(curr_run)
        if curr_run != "":
            print(curr_run)

# titles = bold_titles(doc) # So far missing a few

def sortEntries(allsongs:dict[str,str]) -> dict[str, str]:
    """Sort the entires by song num
    """
    # Returns sorted keys
    sorted_keys = sorted(allsongs.keys(), key=int)
    # Adding the songs in new sorted order
    sorted_dict = {}
    for key in sorted_keys:
        sorted_dict[key] = allsongs[key]
    return sorted_dict

song_nums_and_titles: dict[str, str] = {}
def titles_from_tsank():
    for line in doc.paragraphs:
        # if findall(r'\t', line)
        if '\t' in line.text:
            try:
                song_num = findall(r'\d.*', line.text)[0]
                line_text = sub(r"[^ա-ֆԱ-Ֆ և \s]", "", line.text).replace('\t', '')
                song_title = findall(r'[ա-ֆԱ-Ֆ և \s]*', line_text)[0]
                song_nums_and_titles[song_num] = song_title
            except:
                print("This line failed?")
                print(line.text)

titles_from_tsank()
# Gets all song titles
song_nums_and_titles = sortEntries(song_nums_and_titles)
# print(song_nums_and_titles)
def export_titles(all_song_titles):
    #TODO: Make this a tsank.json
    with open("utils/song_nums_and_titles.json", 'w', encoding='utf-8') as file:
        dump(all_song_titles, file, indent=4, ensure_ascii=False)

export_titles(song_nums_and_titles)

class IncrementedList():
    def __init__(self, curr_list) -> None:
        if isinstance(curr_list, list):
            self.list = curr_list
        else:
            # Attempt conversion
            self.list = list(curr_list) #eg: [('123','Ser e Astvac'), ...]
        self.curr_index = 0
        self.len= len(self.list)

    def get_current_value(self): return self.list[self.curr_index]
    def __next__(self) -> tuple[str,str]:
        if len(self.list) != self.curr_index:
            self.curr_index += 1
            return self.list[self.curr_index-1]
        else:
            raise StopIteration("reached end of list")
    def increment_index(self, increment=1): self.curr_index += increment
    def __len__(self) -> int:
        return len(self.list)

song_found = False
title_found = False
all_songs = IncrementedList(song_nums_and_titles.items())
song_docs:dict[str,document.Document] = {}
# for paragrapgh in doc.paragraphs:
#     # Check if the line matches a title.
#     # IF so, close the previous file,
#     # Starting here open a new one

#     # Clears out all none letter chars
#     line_text = sub(r"[^ա-ֆԱ-Ֆ և \s]", "", paragrapgh.text).replace('\t', '')
#     # Searches for all letter chars
#     try:
#         song_title = findall(r'[ա-ֆԱ-Ֆ և \s]*', line_text)[0]
#     except:
#         song_title = ""
#     if song_title == all_titles.get_current_value() and not title_found: title_found = True
#     curr_run = ""
#     for run in paragrapgh.runs:
#         if run.bold and title_found:
#             song_found = True


#     if song_found:
#         print(paragrapgh.text)

# Algo:
# Song starts when you match a line with the current title and that line is bold
# Once that happens everything until that condition is met again is the song.
song_start = False
current_song: tuple[str,str] = next(all_songs) # (Song Num, Title)
current_title: str = current_song[1]
current_songnum: str = current_song[0]
title_line = False # For avoiding adding the title to the doc
previous_line = '' # Used to get rid of all extra empty lines
songnum = ''
EndOfList = False # Used to know when to exit, but still need to make the doc
for line in doc.paragraphs:
    current_line = sub(pattern=r"[^ա-ֆԱ-Ֆ-և\s]", repl="",string=line.text.strip(),count=0,flags=MULTILINE)
    if (current_title in current_line and line.runs[0].bold and not EndOfList):

        # next_song: tuple[str,str] = next(all_songs)
        # next_title: str = next_song[1]
        # next_songnum: str = next_song[0]

        song_start = True
        title_line = True
        song_docs[current_songnum] = Document()
        song_docs[current_songnum].add_paragraph(current_songnum)
        songnum = current_songnum
        try:
            current_song: tuple[str,str] = next(all_songs)
        except StopIteration:
            # break
            EndOfList = True
        current_title: str = current_song[1]
        current_songnum: str = current_song[0]

    if song_start and not title_line:
        if not(previous_line == '' and previous_line == line.text):
            previous_line = line.text
            # Get indentation
            format: ParagraphFormat = line.paragraph_format

            first_line_indent = format.first_line_indent
            left_indent = format.left_indent
            right_indent = format.right_indent
            new_par = song_docs[songnum].add_paragraph(line.text)#, line.style)
            style = song_docs[songnum].styles['Normal']
            font = style.font # type: ignore
            font.name = 'Arial'
            font.size = Pt(22)
            # Set proper indentation
            new_par.paragraph_format.first_line_indent = format.first_line_indent
            new_par.paragraph_format.left_indent = format.left_indent
            new_par.paragraph_format.right_indent = format.right_indent
            # new_par.paragraph_format.line_spacing_rule = format.line_spacing_rule
            # new_par.paragraph_format.line_spacing = format.line_spacing

            if format.space_after:
                # Gave up on the space_after = format.space_after
                # And did it the crrect way with a newline
                song_docs[songnum].add_paragraph()
            new_par.paragraph_format.space_after = 0 # By default it has a line spacing so we set it to 0
            # new_par.paragraph_format.space_after = format.space_before

    if title_line:
        title_line = False
# print(song_docs[current_songnum])
for song, songdoc in song_docs.items():
    songdoc.save(f"utils/songs/{song}.docx")