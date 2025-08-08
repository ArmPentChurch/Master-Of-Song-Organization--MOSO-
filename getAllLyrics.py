from time import time as today
from json import load
from os import environ as ENV
from docx import Document, document
# from docx.document import Document
from SPOT import ALL_LYRICS, DATABASE_FILEPATH
#open all files, get latest version and store lyrics in dict
# if songLyrics.get('latestChange',None) or do < today()

def getAllLyricsDict() -> dict:
    with open(ALL_LYRICS, 'r', encoding='utf-8') as f:
        return load(f)

def GetAllSongPths(index:dict, songs_to_open:dict) -> dict:
    """Fills a dict with song nums to open, with file paths.

    Args:
        index (dict): The song database
        songs_to_open (dict): A dict of songs to open

    Returns:
        dict: A dict of songs to open
    """
    for songNum in index:
        songPth: str = index[songNum].get('latestVersion', None)#['latestVersion']
        if not songPth:
            print(f"This song does not have a latest version {songNum}")
        else:
            songs_to_open[songNum] = songPth
    return songs_to_open

# def readLyrics(filePth:str) -> str:
#     lyrics = ''
#     doc = Document(filePth)
#     for p in doc.paragraphs:
#         lyrics += p.text

#     if lyrics != '':
#         return lyrics
#     return None

def readLyrics(doc:document.Document) -> str:
    lyrics = ''
    # If doc is not a document.Document, make it one.
    if not isinstance(doc, document.Document) and isinstance(doc, str):
        doc = Document(doc)
    # doc = Document(filePth)
    for p in doc.paragraphs:
        lyrics += p.text

    if lyrics != '':
        return lyrics
    return None

def updateSongLyrics(songNum:str, lyrics:document.Document):
    try:
        allLyrics = getAllLyricsDict() # Open lyrics dict
        allLyrics[songNum] = readLyrics(lyrics) # Update specified value
        # Save updated values
        with open(ALL_LYRICS, 'w', encoding='utf-8') as f:
            from json import dump
            dump(obj=allLyrics, fp=f, ensure_ascii=False, indent=4)
        return True
    except Exception as e:
        print(f"Something went wrong:\n{e}")
        return False

def getAllLyrics():
    songs_to_open = {}
    songLyrics = {
        'latestChange': today()
    }
    with open(DATABASE_FILEPATH, 'r', encoding='utf-8') as f:
        all_songs = load(f)
        GetAllSongPths(all_songs,songs_to_open)
    for songNum, filePth in songs_to_open.items():
        songNum:str
        lyrics = readLyrics(filePth)
        songLyrics[songNum] = lyrics
    with open(ALL_LYRICS, 'w', encoding='utf-8') as f:
        from json import dump
        dump(obj=songLyrics, fp=f, ensure_ascii=False, indent=4)

if __name__ == "__main__":
    pass
    # updateSongLyrics('664')
    # singleWordToJson('664')
    # singleWordToJson('95')
    # singleWordToJson('96')
    # onedrive = ENV.get('onedrive')
    # updateSongLyrics('old','389',Document(onedrive+'\\Word songs/389 Տոն է այսոր սուրբ հաղթական.docx'))
